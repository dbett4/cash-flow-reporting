#!/usr/bin/env python3
"""Author only the fictional source workbooks and their exact-coordinate manifest.

This is an explicit fixture maintenance command, never part of report generation.
It does not produce or update the independently frozen expected.json.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
CASE = ROOT / "examples/harbor"
ENTITIES = ("parent", "subsidiary")
COLS = [*ENTITIES, "eliminations", "consolidated"]


def fact(name):
    return {"kind": "fact", "id": name}


def total(*terms):
    return {"kind": "sum", "terms": list(terms)}


def zero(reason="No activity under the fictional case policy"):
    return {"kind": "zero", "reason": reason}


def subtract(left, right):
    return {"kind": "difference", "left": left, "right": right}


def negative(expr):
    return subtract(zero(), expr)


def ref(section, row, column="consolidated"):
    return {"kind": "reference", "section": section, "row": row, "column": column}


def group(key):
    return total(*(fact(f"{entity}_{key}") for entity in ENTITIES))


def main():
    data = json.loads((CASE / "source-data.json").read_text())
    manifest = {"schema": "acfr.report_manifest.v1", "fiscal_year": 2025,
                "sources": [], "facts": [], "columns": [
                    {"id": "parent", "label": "Parent"},
                    {"id": "subsidiary", "label": "Subsidiary"},
                    {"id": "eliminations", "label": "Eliminations"},
                    {"id": "consolidated", "label": "Consolidated"}],
                "sections": [], "checks": [], "source_reconciliations": []}
    source_sets = dict(data["entities"])
    source_sets["bank"] = data["separate_bank_control"]
    for entity, values in source_sets.items():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Source"
        sheet.append(["Fiscal year ended June 30", 2025])
        sheet.append(["Period start", data["period"]["start"]])
        sheet.append(["Period end", data["period"]["end"]])
        sheet.append(["Entity / control", values.get("label", "Separate fictional bank control")])
        for key, amount in values.items():
            if key == "label":
                continue
            caption = key.replace("_", " ").capitalize()
            sheet.append([caption, amount])
            cell = sheet.max_row
            manifest["facts"].append({"id": f"{entity}_{key}", "source": entity,
                "cell": f"B{cell}", "anchors": {f"A{cell}": caption,
                "A1": "Fiscal year ended June 30", "B2": "2024-07-01", "B3": "2025-06-30",
                "B4": values.get("label", "Separate fictional bank control")}})
            sheet.cell(cell, 2).number_format = '#,##0.00;(#,##0.00);"-"'
        sheet.column_dimensions["A"].width = 39
        sheet.column_dimensions["B"].width = 29
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        path = CASE / f"{entity}.xlsx"
        workbook.save(path)
        manifest["sources"].append({"id": entity, "path": path.name,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "sheet": "Source", "period_cell": "B1"})

    def add_section(section_id, title, definitions, consolidated_only=False):
        rows = []
        for row_id, label, expressions in definitions:
            if consolidated_only:
                values = {"consolidated": expressions}
            else:
                values = dict(zip(COLS[:3], expressions))
                values["consolidated"] = total(*(ref(section_id, row_id, col) for col in COLS[:3]))
            rows.append({"id": row_id, "label": label, "decimals": 2, "values": values})
        section = {"id": section_id, "title": title, "rows": rows}
        if consolidated_only:
            section["columns"] = ["consolidated"]
        manifest["sections"].append(section)

    def movement(key, sign=1):
        return [(negative if sign < 0 else lambda x: x)(fact(f"{e}_{key}")) for e in ENTITIES] + [zero()]

    def section_sum(row_ids):
        return [total(*(ref("cash_flows", r, c) for r in row_ids)) for c in COLS[:3]]

    lending = movement("intercompany_loan_advanced", -1)
    lending[2] = fact("parent_intercompany_loan_advanced")
    borrowing = movement("intercompany_loan_received")
    borrowing[2] = negative(fact("subsidiary_intercompany_loan_received"))
    beginning = [total(fact(f"{e}_cash_beginning"), fact(f"{e}_restricted_cash_beginning")) for e in ENTITIES] + [zero()]
    add_section("cash_flows", "Statement of cash flows — indirect method (USD)", [
        ("net_income", "Net income", movement("net_income")),
        ("depreciation", "Depreciation", movement("depreciation")),
        ("receivables", "Increase in trade receivables", movement("receivables_increase", -1)),
        ("inventory", "Increase in inventory", movement("inventory_increase", -1)),
        ("payables", "Increase in trade payables", movement("payables_increase")),
        ("operating", "Net cash provided by operating activities", section_sum(["net_income", "depreciation", "receivables", "inventory", "payables"])),
        ("equipment", "Cash purchases of equipment", movement("cash_equipment_purchases", -1)),
        ("internal_lending", "Loan advanced to subsidiary (eliminated)", lending),
        ("investing", "Net cash used in investing activities", section_sum(["equipment", "internal_lending"])),
        ("debt_proceeds", "Proceeds from external borrowing", movement("external_debt_proceeds")),
        ("debt_repayments", "Principal repayments of external debt", movement("external_debt_repayments", -1)),
        ("dividends", "Dividends paid to external shareholders", movement("dividends_paid", -1)),
        ("internal_borrowing", "Loan received from parent (eliminated)", borrowing),
        ("financing", "Net cash provided by financing activities", section_sum(["debt_proceeds", "debt_repayments", "dividends", "internal_borrowing"])),
        ("net_change", "Increase in cash and restricted cash", section_sum(["operating", "investing", "financing"])),
        ("beginning", "Cash and restricted cash — beginning", beginning),
        ("ending", "Cash and restricted cash — ending", section_sum(["beginning", "net_change"]))])
    # Consolidated subtotals consume the consolidated rows. Summing only the
    # entity subtotals could mask an omitted entity in a preceding report row.
    consolidated_subtotals = {
        "operating": ["net_income", "depreciation", "receivables", "inventory", "payables"],
        "investing": ["equipment", "internal_lending"],
        "financing": ["debt_proceeds", "debt_repayments", "dividends", "internal_borrowing"],
        "net_change": ["operating", "investing", "financing"],
        "ending": ["beginning", "net_change"],
    }
    for row in manifest["sections"][0]["rows"]:
        if row["id"] in consolidated_subtotals:
            row["values"]["consolidated"] = total(*(ref("cash_flows", r) for r in consolidated_subtotals[row["id"]]))
    add_section("cash_components", "Cash and restricted cash reconciliation (USD)", [
        ("cash_beginning", "Unrestricted cash — beginning", group("cash_beginning")),
        ("restricted_beginning", "Restricted cash — beginning", group("restricted_cash_beginning")),
        ("total_beginning", "Combined cash — beginning", total(ref("cash_components", "cash_beginning"), ref("cash_components", "restricted_beginning"))),
        ("cash_ending", "Unrestricted cash — ending", group("cash_ending")),
        ("restricted_ending", "Restricted cash — ending", group("restricted_cash_ending")),
        ("total_ending", "Combined cash — ending", total(ref("cash_components", "cash_ending"), ref("cash_components", "restricted_ending"))),
        ("cash_change", "Increase in unrestricted cash", subtract(ref("cash_components", "cash_ending"), ref("cash_components", "cash_beginning"))),
        ("restricted_change", "Increase in restricted cash", subtract(ref("cash_components", "restricted_ending"), ref("cash_components", "restricted_beginning"))),
        ("combined_change", "Combined cash increase (transfers excluded from activities)", total(ref("cash_components", "cash_change"), ref("cash_components", "restricted_change"))),
        ("bank_confirmation", "Separate fictional bank control — combined cash ending", fact("bank_cash_and_restricted_cash_ending")),
        ("bank_difference", "Unresolved difference: ledger less bank control", subtract(ref("cash_components", "total_ending"), ref("cash_components", "bank_confirmation")))], True)
    add_section("supplemental", "Supplemental cash flow disclosures (USD)", [
        ("noncash_equipment", "Equipment acquired directly by issuing seller note (no cash)", group("noncash_equipment_seller_note")),
        ("interest_paid", "Cash interest paid (included in operating cash flows)", group("cash_interest_paid")),
        ("taxes_paid", "Cash income taxes paid (included in operating cash flows)", group("cash_income_taxes_paid"))], True)
    add_section("external_debt", "External debt rollforward — intercompany debt excluded (USD)", [
        ("beginning", "External debt — beginning", group("external_debt_beginning")),
        ("cash_proceeds", "External cash borrowings", ref("cash_flows", "debt_proceeds")),
        ("cash_repayments", "External principal repayments", ref("cash_flows", "debt_repayments")),
        ("seller_note", "Noncash seller note issued for equipment", ref("supplemental", "noncash_equipment")),
        ("ending", "External debt — ending", total(*(ref("external_debt", r) for r in ["beginning", "cash_proceeds", "cash_repayments", "seller_note"])))], True)
    add_section("equipment_rollforward", "Net equipment rollforward (USD)", [
        ("beginning", "Net equipment — beginning", group("ppe_net_beginning")),
        ("cash_purchases", "Cash equipment purchases", negative(ref("cash_flows", "equipment"))),
        ("noncash_purchases", "Noncash equipment purchased with seller note", ref("supplemental", "noncash_equipment")),
        ("depreciation", "Depreciation", negative(ref("cash_flows", "depreciation"))),
        ("ending", "Net equipment — ending", total(*(ref("equipment_rollforward", r) for r in ["beginning", "cash_purchases", "noncash_purchases", "depreciation"])))], True)
    checks = manifest["checks"]
    for entity in ENTITIES:
        checks.append({"id": f"{entity}_cash_bridge", "left": ref("cash_flows", "ending", entity),
            "right": total(fact(f"{entity}_cash_ending"), fact(f"{entity}_restricted_cash_ending"))})
    checks.extend([
        {"id": "intercompany_legs_match", "left": fact("parent_intercompany_loan_advanced"), "right": fact("subsidiary_intercompany_loan_received")},
        {"id": "internal_lending_eliminated", "left": ref("cash_flows", "internal_lending"), "right": zero("Intragroup lending excluded from consolidated cash flows")},
        {"id": "internal_borrowing_eliminated", "left": ref("cash_flows", "internal_borrowing"), "right": zero("Intragroup borrowing excluded from consolidated cash flows")},
        {"id": "cash_only_equipment", "left": ref("cash_flows", "equipment"), "right": negative(group("cash_equipment_purchases"))},
        {"id": "cash_only_borrowings", "left": ref("cash_flows", "debt_proceeds"), "right": group("external_debt_proceeds")},
        {"id": "combined_opening", "left": ref("cash_flows", "beginning"), "right": ref("cash_components", "total_beginning")},
        {"id": "combined_closing", "left": ref("cash_flows", "ending"), "right": ref("cash_components", "total_ending")},
        {"id": "combined_movement", "left": ref("cash_flows", "net_change"), "right": ref("cash_components", "combined_change")},
        {"id": "external_debt_bridge", "left": ref("external_debt", "ending"), "right": group("external_debt_ending")},
        {"id": "equipment_bridge", "left": ref("equipment_rollforward", "ending"), "right": group("ppe_net_ending")}])
    manifest["source_reconciliations"].append({"id": "ledger_to_bank", "left": ref("cash_components", "total_ending"), "right": ref("cash_components", "bank_confirmation")})
    (CASE / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print("Authored three fictional workbooks and manifest; expected.json unchanged.")


if __name__ == "__main__":
    main()
