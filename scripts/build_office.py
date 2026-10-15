#!/usr/bin/env python3
"""Compile source-backed report sections, then build editable Office artifacts.

Rust owns source decoding, exact arithmetic and dependency compilation. This
wrapper materializes numeric formulas into local XLSX and DOCX artifacts. Local
arithmetic and Office construction never imply accounting approval.
"""
from __future__ import annotations

import argparse
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import time
import xml.etree.ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parents[1]
NUMBER_FORMAT = '#,##0;(#,##0);"-"'


def cache_compiled_amounts(path: Path, plan: dict) -> None:
    """Keep formulas and cache their same-generation exact compiler predictions.

    This is numeric OOXML data, not a second presentation sheet or a claim that
    Excel recalculated the workbook. Source workbook formulas are not evaluated.
    """
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    for index, section in enumerate(plan["sections"], 2):
        name = f"xl/worksheets/sheet{index}.xml"
        tree = ET.fromstring(entries[name])
        columns = section.get("columns", plan["columns"])
        predictions = {
            f"{chr(66 + c)}{r}": row["cells"][column["id"]]
            for r, row in enumerate(section["rows"], 5)
            for c, column in enumerate(columns)
        }
        for cell in tree.findall(".//s:c", ns):
            predicted = predictions.get(cell.get("r"))
            if predicted is None:
                continue
            formula = cell.find("s:f", ns)
            if formula is None or formula.text != predicted["formula"]:
                raise ValueError("formula changed before numeric cache binding")
            value = cell.find("s:v", ns)
            if value is None:
                value = ET.SubElement(cell, "{" + ns["s"] + "}v")
            value.text = predicted["amount"]
            cell.set("t", "n")
        entries[name] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def compile_report(manifest: Path, output: Path, compiler: Path) -> dict:
    manifest = manifest.resolve()
    output.mkdir(parents=True, exist_ok=True)
    plan_path = output / "report-plan.json"
    started = time.monotonic()
    result = subprocess.run([str(compiler.resolve()), str(manifest), str(plan_path)],
                            capture_output=True, text=True, check=False)
    if result.returncode:
        raise ValueError(result.stderr.strip() or "report compilation failed")
    plan = json.loads(plan_path.read_text())
    if (plan.get("schema") != "acfr.report_plan.v1"
            or plan.get("manifest_sha256") != digest(manifest)
            or plan.get("local_arithmetic_passed") is not True
            or not plan.get("checks")
            or any(check.get("passed") is not True
                   or not 0 <= Decimal(check.get("tolerance", "0")) <= Decimal("0.000001")
                   or abs(Decimal(check["residual"])) > Decimal(check.get("tolerance", "0"))
                   or (Decimal(check.get("tolerance", "0")) and not check.get("residual_reason"))
                   for check in plan["checks"])
            or plan.get("native_verified") is not False
            or plan.get("client_ready") is not False):
        raise ValueError("compiler plan does not match the current admitted manifest")
    plan["compile_seconds"] = time.monotonic() - started
    plan["compiler_sha256"] = digest(compiler.resolve())
    # Bind the actual plan consumed by the Office builder, including its tool
    # identity, rather than leaving a different plan on disk.
    plan_path.write_text(json.dumps(plan, indent=2) + "\n")
    return plan


def _display(value: str, decimals: int = 0) -> str:
    # Only the temporary DOCX's starting text uses display rounding. Numeric
    # spreadsheet values and formulas remain unscaled and unrounded.
    amount = Decimal(value).quantize(Decimal(10) ** -decimals, rounding=ROUND_HALF_UP)
    if not amount:
        return "-"
    return f"({abs(amount):,f})" if amount < 0 else f"{amount:,f}"


def build_artifacts(plan: dict, output: Path, *, title: str, narrative: str = "") -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt

    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    facts = workbook.active
    facts.title = "Source Facts"
    facts.append(["Fact control", "Source", "Sheet", "Cell", "Source SHA-256", "Amount", "Cached source formula"])
    for fact in plan["facts"]:
        facts.append([fact["id"], fact["source"], fact["sheet"], fact["cell"],
                      fact["source_sha256"], Decimal(fact["amount"]), fact.get("cached_formula")])
        # Source formula text is retained as provenance, never executed against
        # nonexistent external workbooks in the generated report.
        if fact.get("cached_formula"):
            facts.cell(facts.max_row, 7).data_type = "s"
        facts.cell(facts.max_row, 6).number_format = NUMBER_FORMAT
        facts.cell(facts.max_row, 6).font = Font(color="0000FF")
    facts.freeze_panes = "F2"
    facts.column_dimensions["A"].width = 35
    facts.column_dimensions["F"].width = 20

    document = Document()
    page = document.sections[0]
    page.orientation = WD_ORIENT.LANDSCAPE
    page.page_width, page.page_height = Inches(11), Inches(8.5)
    page.top_margin = page.bottom_margin = Inches(0.45)
    page.left_margin = page.right_margin = Inches(0.4)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(8)
    normal.paragraph_format.space_after = Pt(0)
    bindings = []
    for index, section in enumerate(plan["sections"]):
        columns = section.get("columns", plan["columns"])
        last_col = get_column_letter(len(columns) + 1)
        if index:
            document.add_page_break()
        document.add_heading(title, level=1)
        document.add_heading(section["title"], level=2)
        document.add_paragraph(f"For the year ended June 30, {plan['fiscal_year']}")
        if index == len(plan["sections"]) - 1 and narrative:
            document.add_paragraph(narrative)
        sheet = workbook.create_sheet(section["id"])
        sheet.cell(1, 1, section["title"])
        sheet.cell(1, 1).data_type = "s"
        sheet.cell(2, 1, plan["fiscal_year"])
        sheet.cell(4, 1, "Description")
        for c, column in enumerate(columns, 2):
            sheet.cell(4, c, column["label"])
            sheet.cell(4, c).data_type = "s"
        table = document.add_table(rows=1, cols=len(columns) + 1)
        table.autofit = False
        table.cell(0, 0).text = "Description"
        for c, column in enumerate(columns, 1):
            table.cell(0, c).text = column["label"]
        for r, row in enumerate(section["rows"], 5):
            sheet.cell(r, 1, row["label"])
            sheet.cell(r, 1).data_type = "s"
            cells = table.add_row().cells
            cells[0].text = row["label"]
            for c, column in enumerate(columns, 2):
                predicted = row["cells"][column["id"]]
                formula = "=" + predicted["formula"]
                if any(token in formula.upper() for token in ("TEXT(", "ROUND(")):
                    raise ValueError("presentation transformation in numeric formula")
                decimals = row.get("decimals", 0)
                digit_pattern = '#,##0' + ('.' + '0' * decimals if decimals else '')
                sheet.cell(r, c, formula).number_format = f'{digit_pattern};({digit_pattern});"-"'
                cells[c - 1].text = _display(predicted["amount"], decimals)
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = Font(name="Times New Roman", size=9)
                cell.alignment = Alignment(vertical="center", wrap_text=cell.column == 1)
        for cell in sheet[4]:
            cell.font = Font(name="Times New Roman", size=9, bold=True)
            cell.fill = PatternFill("solid", fgColor="E9EEF2")
        sheet.column_dimensions["A"].width = 75
        for c in range(2, len(columns) + 2):
            sheet.column_dimensions[get_column_letter(c)].width = 19
        sheet.freeze_panes = "B5"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        for row in table.rows:
            for c, cell in enumerate(row.cells):
                label_width = 4.2 if len(columns) <= 5 else 3.5
                cell.width = Inches(label_width if c == 0 else (10.2 - label_width) / len(columns))
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name, run.font.size = "Times New Roman", Pt(8)
        bindings.append({"section": section["id"], "table_index": index,
                         "source_range": f"A4:{last_col}{len(section['rows']) + 4}",
                         "rows": len(section["rows"]) + 1, "columns": len(columns) + 1})
    xlsx, docx = output / "report.xlsx", output / "report.docx"
    workbook.save(xlsx)
    cache_compiled_amounts(xlsx, plan)
    document.save(docx)
    plan_path = output / "report-plan.json"
    if json.loads(plan_path.read_text()) != plan:
        raise ValueError("Office builder plan differs from the persisted compilation")
    record = {"schema": "acfr.report_artifacts.v1", "manifest_sha256": plan["manifest_sha256"],
              "plan_sha256": digest(plan_path), "compiler_sha256": plan["compiler_sha256"],
              "builder_sha256": digest(Path(__file__)),
              "xlsx": str(xlsx.resolve()), "xlsx_sha256": digest(xlsx),
              "docx": str(docx.resolve()), "docx_sha256": digest(docx), "bindings": bindings,
              "compile_seconds": plan["compile_seconds"], "native_verified": False, "client_ready": False}
    generation = {key: record[key] for key in ("manifest_sha256", "plan_sha256", "compiler_sha256",
                                               "builder_sha256", "xlsx_sha256", "docx_sha256", "bindings")}
    record["generation_sha256"] = hashlib.sha256(json.dumps(generation, sort_keys=True,
                                                           separators=(",", ":")).encode()).hexdigest()
    (output / "report-artifacts.json").write_text(json.dumps(record, indent=2) + "\n")
    return record


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("manifest", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--compiler", type=Path,
                        default=ROOT / "tools/report-compiler/target/release/acfr-report-compiler")
    parser.add_argument("--title", required=True)
    parser.add_argument("--narrative", default="")
    args = parser.parse_args(argv)
    try:
        plan = compile_report(args.manifest, args.out, args.compiler)
        record = build_artifacts(plan, args.out, title=args.title, narrative=args.narrative)
        print(json.dumps({"sections": len(record["bindings"]), "source_facts": len(plan["facts"]),
                          "accounting_checks": len(plan["checks"]), "compile_seconds": record["compile_seconds"],
                          "artifacts": str((args.out / "report-artifacts.json").resolve()),
                          "native_verified": False, "client_ready": False}))
        return 0
    except (ValueError, OSError, KeyError) as exc:
        print(str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
