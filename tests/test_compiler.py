from __future__ import annotations

import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
BINARY = ROOT / "tools/report-compiler/target/release/acfr-report-compiler"


def reference(section, row):
    return {"kind": "reference", "section": section, "row": row, "column": "fund"}


def script_module(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / f"scripts/{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class ReportCompilerTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.addCleanup(self.temp.cleanup)
        self.source = self.root / "source.xlsx"
        workbook = Workbook()
        workbook.active.title = "Source"
        for value in (2025, 100, 0.1, 0.2, 100.3):
            workbook.active.append([value])
        workbook.save(self.source)
        fact = lambda name: {"kind": "fact", "id": name}
        self.manifest = {
            "schema": "acfr.report_manifest.v1", "fiscal_year": 2025,
            "sources": [{"id": "source", "path": "source.xlsx", "sheet": "Source", "period_cell": "A1",
                         "sha256": hashlib.sha256(self.source.read_bytes()).hexdigest()}],
            "facts": [{"id": name, "source": "source", "cell": cell} for name, cell in
                      (("begin", "A2"), ("activity_a", "A3"), ("activity_b", "A4"), ("end", "A5"))],
            "columns": [{"id": "fund", "label": "Fund"}],
            "sections": [
                {"id": "statement", "title": "Statement", "rows": [
                    {"id": "begin", "label": "Opening", "values": {"fund": fact("begin")}},
                    {"id": "activity", "label": "Activity", "values": {"fund": {"kind": "sum", "terms": [fact("activity_a"), fact("activity_b")]}}},
                    {"id": "end", "label": "Closing", "values": {"fund": {"kind": "sum", "terms": [reference("statement", "begin"), reference("statement", "activity")]}}},
                ]},
                {"id": "note", "title": "Note", "rows": [
                    {"id": "closing", "label": "Disclosed closing", "values": {"fund": reference("statement", "end")}},
                ]},
            ],
            "checks": [{"id": "closing_control", "left": reference("note", "closing"), "right": fact("end")}],
        }

    def compile(self):
        manifest = self.root / "manifest.json"
        manifest.write_text(json.dumps(self.manifest))
        result = subprocess.run([str(BINARY), str(manifest), str(self.root / "plan.json")], capture_output=True, text=True)
        return result, json.loads((self.root / "plan.json").read_text()) if (self.root / "plan.json").exists() else None

    def test_source_to_statement_to_note_preserves_exact_decimal_control(self):
        result, plan = self.compile()
        self.assertEqual(result.returncode, 0, result.stderr)
        cell = plan["sections"][1]["rows"][0]["cells"]["fund"]
        self.assertEqual(cell["amount"], "100.3")
        self.assertEqual(cell["formula"], "'statement'!B7")
        self.assertEqual(set(cell["fact_ids"]), {"begin", "activity_a", "activity_b"})
        self.assertFalse(plan["native_verified"])
        self.assertFalse(plan["client_ready"])

    def test_nonzero_financial_residual_is_retained_and_refuses_admission(self):
        self.manifest["checks"][0]["right"] = {"kind": "fact", "id": "begin"}
        result, plan = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(plan["checks"][0]["residual"], "0.3")
        self.assertFalse(plan["local_arithmetic_passed"])

    def test_source_hash_change_refuses_replay(self):
        self.manifest["sources"][0]["sha256"] = "0" * 64
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("source hash mismatch", result.stderr)

    def test_source_period_change_refuses_replay(self):
        self.manifest["fiscal_year"] = 2026
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("source period mismatch", result.stderr)

    def test_cycle_is_rejected(self):
        self.manifest["sections"][0]["rows"][2]["values"]["fund"] = reference("note", "closing")
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("cyclic report dependency", result.stderr)

    def test_missing_column_cannot_become_an_implicit_zero(self):
        self.manifest["columns"].append({"id": "another", "label": "Another"})
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("incomplete column coverage", result.stderr)

    def test_material_tolerance_cannot_hide_a_residual(self):
        self.manifest["checks"][0].update(tolerance="1", residual_reason="ignore it")
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("sub-cent", result.stderr)

    def test_generated_faces_are_formulas_and_literal_labels_remain_text(self):
        module = script_module("build_office")
        self.manifest["sections"][0]["rows"][0]["label"] = "=untrusted label"
        manifest = self.root / "manifest.json"
        manifest.write_text(json.dumps(self.manifest))
        plan = module.compile_report(manifest, self.root, BINARY)
        record = module.build_artifacts(plan, self.root, title="Synthetic report")
        workbook = load_workbook(record["xlsx"], data_only=False)
        self.assertEqual(workbook["statement"]["A5"].data_type, "s")
        self.assertEqual(workbook["statement"]["B7"].data_type, "f")
        self.assertEqual(workbook["note"]["B5"].value, "='statement'!B7")
        self.assertEqual(workbook["Source Facts"]["F2"].value, 100)
        workbook.close()

    def artifacts(self):
        builder = script_module("build_office")
        manifest = self.root / "manifest.json"
        manifest.write_text(json.dumps(self.manifest))
        plan = builder.compile_report(manifest, self.root, BINARY)
        record = builder.build_artifacts(plan, self.root, title="Synthetic report")
        return builder, plan, record





    def test_changed_source_caption_refuses_reused_mapping(self):
        workbook = load_workbook(self.source)
        workbook.active['B2'] = 'Opening cash'
        workbook.save(self.source)
        self.manifest['sources'][0]['sha256'] = hashlib.sha256(self.source.read_bytes()).hexdigest()
        self.manifest['facts'][0]['anchors'] = {'B2': 'Operating receipts'}
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('fact caption changed', result.stderr)

    def test_references_use_the_selected_section_column_position(self):
        self.manifest['columns'].insert(0, {'id': 'unused', 'label': 'Unused'})
        for section in self.manifest['sections']:
            section['columns'] = ['fund']
        result, plan = self.compile()
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(plan['sections'][1]['rows'][0]['cells']['fund']['formula'], "'statement'!B7")
        self.assertEqual([c['id'] for c in plan['sections'][0]['columns']], ['fund'])

    def test_source_conflict_remains_visible_without_disabling_arithmetic(self):
        self.manifest['source_reconciliations'] = [{'id': 'source_difference',
            'left': reference('note', 'closing'), 'right': {'kind': 'fact', 'id': 'begin'}}]
        result, plan = self.compile()
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(plan['local_arithmetic_passed'])
        self.assertFalse(plan['source_reconciliations'][0]['passed'])
        self.assertEqual(plan['source_reconciliations'][0]['residual'], '0.3')
        self.assertFalse(plan['client_ready'])

    def test_cents_use_native_format_without_rounding_the_formula(self):
        from docx import Document
        self.manifest['sections'][0]['rows'][1]['decimals'] = 2
        _, plan, record = self.artifacts()
        workbook = load_workbook(record['xlsx'], data_only=False)
        cell = workbook['statement']['B6']
        self.assertEqual(cell.number_format, '#,##0.00;(#,##0.00);"-"')
        self.assertEqual(cell.data_type, 'f')
        self.assertNotIn('ROUND', cell.value)
        self.assertEqual(plan['sections'][0]['rows'][1]['cells']['fund']['amount'], '0.3')
        self.assertEqual(Document(record['docx']).tables[0].cell(2, 1).text, '0.30')
        workbook.close()

    def test_unbounded_presentation_precision_is_rejected(self):
        self.manifest['sections'][0]['rows'][1]['decimals'] = 7
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('six decimals', result.stderr)


if __name__ == "__main__":
    unittest.main()
