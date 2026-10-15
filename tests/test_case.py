from __future__ import annotations

from copy import deepcopy
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import demo

BINARY = ROOT / "tools/report-compiler/target/release/acfr-report-compiler"


class CorporateCashFlowTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.case = self.root / "case"
        shutil.copytree(ROOT / "examples/harbor", self.case)
        self.manifest = json.loads((self.case / "manifest.json").read_text())
        self.out = self.root / "output"
        self.out.mkdir()

    def compile(self):
        manifest = self.case / "manifest.json"
        manifest.write_text(json.dumps(self.manifest))
        result = subprocess.run([str(BINARY), str(manifest), str(self.out / "plan.json")], capture_output=True, text=True)
        plan = json.loads((self.out / "plan.json").read_text()) if (self.out / "plan.json").exists() else None
        return result, plan

    def row(self, section, row):
        return next(r for s in self.manifest["sections"] if s["id"] == section for r in s["rows"] if r["id"] == row)

    def source_change(self, source_id, cell, value):
        source = next(s for s in self.manifest["sources"] if s["id"] == source_id)
        path = self.case / source["path"]
        workbook = load_workbook(path)
        workbook["Source"][cell] = value
        workbook.save(path)
        workbook.close()
        source["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()

    def test_frozen_expected_values_cover_every_output_cell(self):
        result, plan = self.compile()
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(demo.check_expected(plan, json.loads((self.case / "expected.json").read_text())), [])
        self.assertEqual(len(plan["checks"]), 12)
        self.assertTrue(all(c["passed"] for c in plan["checks"]))
        self.assertTrue(all(c["passed"] for c in plan["source_reconciliations"]))

    def test_duplicate_source_identity_is_rejected(self):
        self.manifest["sources"].append(deepcopy(self.manifest["sources"][0]))
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("duplicate source id", result.stderr)

    def test_omitted_entity_is_not_an_implicit_zero(self):
        self.manifest["sources"] = [s for s in self.manifest["sources"] if s["id"] != "subsidiary"]
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("unknown fact source", result.stderr)

    def test_subsidiary_omitted_only_from_consolidation_fails_closing(self):
        row = self.row("cash_flows", "operating")
        row["values"]["consolidated"] = {"kind": "sum", "terms": [
            {"kind": "reference", "section": "cash_flows", "row": "operating", "column": "parent"},
            {"kind": "reference", "section": "cash_flows", "row": "operating", "column": "eliminations"}]}
        result, plan = self.compile()
        self.assertNotEqual(result.returncode, 0)
        controls = {c["id"]: c for c in plan["checks"]}
        self.assertFalse(controls["combined_closing"]["passed"])
        self.assertEqual(Decimal(controls["combined_closing"]["residual"]), Decimal("-50000"))

    def test_duplicated_cash_movement_fails_entity_bridge(self):
        row = self.row("cash_flows", "investing")
        row["values"]["parent"]["terms"].append(deepcopy(row["values"]["parent"]["terms"][0]))
        result, plan = self.compile()
        self.assertNotEqual(result.returncode, 0)
        controls = {c["id"]: c for c in plan["checks"]}
        self.assertEqual(Decimal(controls["parent_cash_bridge"]["residual"]), Decimal("-60000"))

    def test_noncash_in_both_cash_sections_fails_even_when_net_change_still_ties(self):
        noncash = {"kind": "fact", "id": "parent_noncash_equipment_seller_note"}
        equipment = self.row("cash_flows", "equipment")["values"]["parent"]
        self.row("cash_flows", "equipment")["values"]["parent"] = {"kind": "difference", "left": equipment, "right": noncash}
        borrowing = self.row("cash_flows", "debt_proceeds")["values"]["parent"]
        self.row("cash_flows", "debt_proceeds")["values"]["parent"] = {"kind": "sum", "terms": [borrowing, noncash]}
        result, plan = self.compile()
        self.assertNotEqual(result.returncode, 0)
        controls = {c["id"]: c for c in plan["checks"]}
        self.assertTrue(controls["combined_closing"]["passed"])
        self.assertFalse(controls["cash_only_equipment"]["passed"])
        self.assertFalse(controls["cash_only_borrowings"]["passed"])

    def test_missing_lending_elimination_fails_despite_offsetting_net_effect(self):
        self.row("cash_flows", "internal_lending")["values"]["eliminations"] = {"kind": "zero", "reason": "Injected omission"}
        self.row("cash_flows", "internal_borrowing")["values"]["eliminations"] = {"kind": "zero", "reason": "Injected omission"}
        result, plan = self.compile()
        self.assertNotEqual(result.returncode, 0)
        controls = {c["id"]: c for c in plan["checks"]}
        self.assertTrue(controls["combined_closing"]["passed"])
        self.assertFalse(controls["internal_lending_eliminated"]["passed"])
        self.assertFalse(controls["internal_borrowing_eliminated"]["passed"])

    def test_wrong_period_is_rejected_after_valid_hash_update(self):
        self.source_change("parent", "B1", 2026)
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("source period mismatch", result.stderr)

    def test_shifted_caption_is_rejected_after_valid_hash_update(self):
        spec = next(f for f in self.manifest["facts"] if f["id"] == "parent_cash_equipment_purchases")
        anchor = next(a for a in spec["anchors"] if a.startswith("A") and a != "A1")
        self.source_change("parent", anchor, "Noncash equipment purchases")
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("fact caption changed", result.stderr)

    def test_wrong_entity_label_is_rejected(self):
        self.source_change("subsidiary", "B4", "Different LLC")
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("fact caption changed", result.stderr)

    def test_wrong_source_hash_is_rejected(self):
        self.manifest["sources"][0]["sha256"] = "0" * 64
        result, _ = self.compile()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("source hash mismatch", result.stderr)

    def test_one_cent_source_difference_is_retained_without_plugging(self):
        spec = next(f for f in self.manifest["facts"] if f["source"] == "bank")
        self.source_change("bank", spec["cell"], 220000.01)
        result, plan = self.compile()
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(plan["local_arithmetic_passed"])
        control = plan["source_reconciliations"][0]
        self.assertFalse(control["passed"])
        self.assertEqual(Decimal(control["residual"]), Decimal("-0.01"))
        self.assertFalse(plan["client_ready"])
        record = demo.build_artifacts(self.persisted(plan), self.out, title="Fictional disagreement")
        workbook = load_workbook(record["xlsx"], data_only=True)
        self.assertEqual(Decimal(str(workbook["cash_components"]["B15"].value)), Decimal("-0.01"))
        workbook.close()

    def persisted(self, plan):
        plan["compile_seconds"] = 0
        plan["compiler_sha256"] = hashlib.sha256(BINARY.read_bytes()).hexdigest()
        (self.out / "report-plan.json").write_text(json.dumps(plan))
        return plan

    def test_real_office_artifacts_preserve_numeric_formulas_and_same_generation_caches(self):
        result, plan = self.compile()
        self.assertEqual(result.returncode, 0, result.stderr)
        record = demo.build_artifacts(self.persisted(plan), self.out, title="Fictional cash flows")
        formulas = load_workbook(record["xlsx"], data_only=False)
        cached = load_workbook(record["xlsx"], data_only=True)
        for section in plan["sections"]:
            for r, row in enumerate(section["rows"], 5):
                for c, col in enumerate(section["columns"], 2):
                    expected = row["cells"][col["id"]]
                    actual = formulas[section["id"]].cell(r, c)
                    self.assertEqual(actual.data_type, "f")
                    self.assertEqual(actual.value, "=" + expected["formula"])
                    self.assertNotIn("ROUND(", actual.value.upper())
                    self.assertNotIn("TEXT(", actual.value.upper())
                    self.assertEqual(Decimal(str(cached[section["id"]].cell(r, c).value)), Decimal(expected["amount"]))
        self.assertEqual(len(Document(record["docx"]).tables), 5)
        self.assertFalse(record["native_verified"])
        formulas.close()
        cached.close()

    def test_demo_returns_nonzero_and_still_builds_reviewable_source_disagreement(self):
        spec = next(f for f in self.manifest["facts"] if f["source"] == "bank")
        self.source_change("bank", spec["cell"], 220000.01)
        (self.case / "manifest.json").write_text(json.dumps(self.manifest))
        result = subprocess.run([sys.executable, str(ROOT / "scripts/demo.py"), "--manifest", str(self.case / "manifest.json"),
            "--expected", str(self.case / "expected.json"), "--out", str(self.out)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 1, result.stderr)
        review = json.loads((self.out / "review.json").read_text())
        self.assertFalse(review["source_reconciliations_passed"])
        self.assertIn("-0.01", (self.out / "review.json").read_text())
        self.assertIn("220,000.01", (self.out / "index.html").read_text())
        self.assertTrue((self.out / "report.xlsx").is_file())
        self.assertTrue((self.out / "report.docx").is_file())

    def test_existing_generation_is_preserved_before_any_new_compilation(self):
        prior = self.out / "report.xlsx"
        prior.write_bytes(b"previous generation is retained")
        result = subprocess.run([sys.executable, str(ROOT / "scripts/demo.py"), "--out", str(self.out)],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertIn("fresh --out", result.stderr)
        self.assertEqual(prior.read_bytes(), b"previous generation is retained")
        self.assertFalse((self.out / "report-plan.json").exists())

    def test_extracted_code_hashes_match_retained_provenance(self):
        provenance = json.loads((ROOT / "SOURCE-PROVENANCE.json").read_text())
        for record in provenance["files"]:
            with self.subTest(path=record["path"]):
                self.assertEqual(hashlib.sha256((ROOT / record["path"]).read_bytes()).hexdigest(),
                                 record["extracted_sha256"])


if __name__ == "__main__":
    unittest.main()
